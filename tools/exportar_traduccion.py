#!/usr/bin/env python3
"""
exportar_traduccion.py — exporta a Excel todo lo que falta traducir.

Una hoja por ficha. En cada hoja, la primera columna es el idioma que se
aprende (el texto original) y las cinco siguientes son las lenguas base. Las
columnas que ya estan completas se marcan en verde y con la cabecera "YA
EXISTE": no hay que tocarlas, y estan ahi solo para que todas las hojas tengan
la misma forma y para dar contexto al traducir.

Lo que NO va en este archivo, porque no es traduccion sino analisis nuevo:
el contraste con la lengua base y los errores contrastivos. El contraste
aleman-italiano dice cosas distintas del aleman-espanol, no las mismas en otro
idioma.

Uso:
    python3 exportar_traduccion.py ../build/*.json --salida ../traducciones.xlsx
"""
import argparse, json, sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

IDIOMAS = ["de", "en", "es", "fr", "it", "pt"]
NOMBRE = {"de": "Alemán", "en": "Inglés", "es": "Español",
          "fr": "Francés", "it": "Italiano", "pt": "Portugués"}

VERDE   = PatternFill("solid", fgColor="D9EAD3")   # ya existe, no tocar
AMBAR   = PatternFill("solid", fgColor="FFF2CC")   # hay que traducir
GRIS    = PatternFill("solid", fgColor="EFEFEF")   # original
CABEZA  = PatternFill("solid", fgColor="1E293B")
SECCION = PatternFill("solid", fgColor="DDE3EA")
BORDE   = Border(*[Side(style="thin", color="BBBBBB")] * 4)


def ancho(ws, anchos):
    for i, a in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = a


def hoja_de_ficha(wb, d):
    aprende = d["idioma"]
    bases = [i for i in IDIOMAS if i != aprende]
    orden = [aprende] + bases

    ws = wb.create_sheet(d["id"][:31])
    ws.freeze_panes = "A5"

    ws["A1"] = f'{d["id"]} · {d["titulo"]}'
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = (f'Idioma que se aprende: {NOMBRE[aprende]} · Nivel: {d["nivel"]} · '
                f'Tema: {d["topicId"]} · {d["subtitulo"]}')
    ws["A2"].font = Font(italic=True, size=9, color="555555")
    ws["A3"] = ("Rellenar solo las columnas en ámbar. Las verdes ya existen y no hay "
                "que traducirlas.")
    ws["A3"].font = Font(size=9, color="7F6000")

    # --- cabecera de columnas ---
    fila = 4
    for j, idi in enumerate(orden, start=1):
        cel = ws.cell(row=fila, column=j)
        if j == 1:
            cel.value = f"{NOMBRE[idi]} — ORIGINAL"
        else:
            cel.value = NOMBRE[idi]
        cel.font = Font(bold=True, color="FFFFFF", size=10)
        cel.fill = CABEZA
        cel.alignment = Alignment(horizontal="center", wrap_text=True)
        cel.border = BORDE
    ancho(ws, [42] + [30] * 5)

    completo = {i: True for i in bases}   # ¿la columna ya está entera?

    def bloque(titulo, filas, campo):
        """filas: lista de (texto_original, dict_de_traducciones)."""
        nonlocal fila
        fila += 1
        c = ws.cell(row=fila, column=1, value=titulo)
        c.font = Font(bold=True, size=10)
        c.fill = SECCION
        for j in range(2, 7):
            ws.cell(row=fila, column=j).fill = SECCION
        for original, trads in filas:
            fila += 1
            c = ws.cell(row=fila, column=1, value=original)
            c.fill = GRIS; c.border = BORDE
            c.alignment = Alignment(wrap_text=True, vertical="top")
            for j, idi in enumerate(orden[1:], start=2):
                cel = ws.cell(row=fila, column=j)
                v = trads.get(idi)
                if idi in trads:
                    # None es deliberado: sin equivalencia directa
                    cel.value = v if v is not None else "—"
                    cel.fill = VERDE
                else:
                    cel.fill = AMBAR
                    completo[idi] = False
                cel.border = BORDE
                cel.alignment = Alignment(wrap_text=True, vertical="top")

    bloque("VOCABULARIO",
           [(v["item"], v["traducciones"]) for v in d["vocabulario"]], "vocabulario")
    bloque("EXPRESIONES",
           [(r["expresion"], r["traducciones"]) for r in d["redemittel"]], "redemittel")

    # Marcar en la cabecera las columnas que NO hay que traducir. La columna
    # se mantiene igual en todas las hojas para que el formato sea uniforme:
    # sirve de contexto al traducir y evita que cada hoja tenga otra forma.
    for j, idi in enumerate(orden[1:], start=2):
        cel = ws.cell(row=4, column=j)
        if completo[idi]:
            cel.value = f"{NOMBRE[idi]} — YA EXISTE, no traducir"
            cel.fill = PatternFill("solid", fgColor="38761D")
        else:
            cel.value = f"{NOMBRE[idi]} — TRADUCIR"
    return ws


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("fichas", nargs="+")
    ap.add_argument("--salida", default="../traducciones.xlsx")
    a = ap.parse_args()

    fichas = []
    for f in sorted(a.fichas):
        d = json.loads(Path(f).read_text(encoding="utf-8"))
        if not d.get("_tipo"):
            fichas.append(d)
    fichas.sort(key=lambda x: (x["idioma"], x["nivel"], x["id"]))

    wb = openpyxl.Workbook()
    idx = wb.active; idx.title = "ÍNDICE"
    idx["A1"] = "Traducciones pendientes"
    idx["A1"].font = Font(bold=True, size=14)
    idx["A2"] = ("Una hoja por ficha. La primera columna es el idioma que se aprende. "
                 "Rellenar solo las celdas en ámbar.")
    idx["A2"].font = Font(size=9, color="555555")
    idx["A3"] = ("NO incluido, porque no es traducción sino análisis nuevo: el contraste "
                 "con la lengua base y los errores contrastivos. Esos dependen del par de "
                 "idiomas y se escriben aparte.")
    idx["A3"].font = Font(size=9, color="7F6000")
    for j, t in enumerate(["Hoja", "Idioma", "Nivel", "Tema", "Título",
                           "Vocab.", "Expr.", "Faltan"], start=1):
        c = idx.cell(row=5, column=j, value=t)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = CABEZA
    ancho(idx, [16, 11, 8, 8, 44, 9, 9, 10])

    total = 0
    for i, d in enumerate(fichas, start=6):
        hoja_de_ficha(wb, d)
        bases = [x for x in IDIOMAS if x != d["idioma"]]
        faltan = sum(1 for v in d["vocabulario"] for b in bases
                     if b not in v["traducciones"])
        faltan += sum(1 for r in d["redemittel"] for b in bases
                      if b not in r["traducciones"])
        total += faltan
        for j, v in enumerate([d["id"], NOMBRE[d["idioma"]], d["nivel"], d["topicId"],
                               d["titulo"], len(d["vocabulario"]),
                               len(d["redemittel"]), faltan], start=1):
            idx.cell(row=i, column=j, value=v)

    fin = len(fichas) + 6
    idx.cell(row=fin, column=5, value="TOTAL").font = Font(bold=True)
    idx.cell(row=fin, column=8, value=total).font = Font(bold=True)

    wb.save(a.salida)
    print(f"{len(fichas)} hojas, {total} celdas por traducir -> {a.salida}",
          file=sys.stderr)


if __name__ == "__main__":
    main()
