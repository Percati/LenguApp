#!/usr/bin/env python3
"""
exportar_traduccion_maestro.py — el maestro único de traducciones pendientes.

Misma filosofia que exportar_audio_maestro.py: una hoja por idioma que se
aprende, agrupado por nivel dentro de cada hoja, formato identico entre
hojas. La diferencia es que aqui cada fila lleva las columnas de destino que
todavia faltan, marcadas con el idioma exacto al que hay que traducir.

Cubre dos cosas:

  1. Vocabulario y expresiones (de nucleos/packs), en una hoja por idioma que
     se aprende.
  2. Toda la prosa de las fichas bilingues A2/B1 (nucleos y apariciones), en
     una hoja por combinacion (idioma + nivel), y la de las semanas de repaso
     A2/B1 (contenido/nucleos/REVIEW-*), en una hoja "Repaso" por combinacion. Segun reglas-fichas.md, en A2
     y B1 la ficha entera se puede mostrar en el idioma de la app, asi que
     cada campo de prosa necesita su traduccion a los otros cinco idiomas.
     Los campos son los 13 que el schema declara bilingues (ver $comment,
     "CAMPOS BILINGUES A2/B1"). La clave del propio idioma de la ficha no se
     pide: es el texto original y la pone el importador.

Deliberadamente NO cubre el contraste ni los errores contrastivos: no son
traduccion, son analisis nuevo por cada par de idiomas, y los escribe el
diseñador aparte.

Uso:
    python3 exportar_traduccion_maestro.py --contenido ../contenido --salida ../traducciones-maestro.xlsx
"""
import argparse, json, re
from pathlib import Path
from collections import defaultdict

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

IDIOMAS = ["de", "en", "es", "fr", "it", "pt"]
NOMBRE = {"de": "Alemán", "en": "Inglés", "es": "Español",
          "fr": "Francés", "it": "Italiano", "pt": "Portugués"}
ORDEN_NIVEL = ["A2", "B1", "B2", "C1", "C2"]
NIVELES_BILINGUES = ["A2", "B1"]

# Limites de longitud que el schema exige a cada campo de prosa (min, max).
# La traduccion tiene que respetarlos igual que el original: si no, la ficha
# no valida al componer. El caso que mas muerde es descripcion, con 200
# caracteres de minimo.
LIMITES = {
    "descripcion": (200, 1600),
    "promptCorreccion": (150, None),
    "notas": (20, None),
    "errores": (10, None),
    "autochequeo": (10, None),
    "ejemplos": (5, None),
    "titulo": (3, 80),
    "subtitulo": (5, 160),
    "mision.consigna": (30, None),
    "microtareas": (15, None),
}


def _limite(ruta):
    raiz = ruta.split("[")[0]
    lim = LIMITES.get(ruta) or LIMITES.get(raiz)
    if not lim:
        return ""
    mn, mx = lim
    return f"{mn}–{mx}" if mx else f"min {mn}"

VERDE  = PatternFill("solid", fgColor="D9EAD3")
AMBAR  = PatternFill("solid", fgColor="FFF2CC")
GRIS   = PatternFill("solid", fgColor="EFEFEF")
CABEZA = PatternFill("solid", fgColor="1E293B")
NIVEL_FILL = PatternFill("solid", fgColor="DDE3EA")
BORDE  = Border(*[Side(style="thin", color="CCCCCC")] * 4)


def limpio(t):
    return re.sub(r"\*+", "", str(t)).strip()


def recolectar(base):
    """{idioma: {nivel: [(tipo, texto, dict_traducciones, origen)]}}, deduplicado."""
    datos = defaultdict(lambda: defaultdict(list))
    vistos = defaultdict(set)

    for p in sorted(base.joinpath("nucleos").glob("*.json")):
        d = json.loads(p.read_text(encoding="utf-8"))
        if d.get("_tipo"):
            continue
        idi, niv, sid = d["idioma"], d["nivel"], d["skillId"]
        for r in d.get("redemittel", []):
            t = limpio(r["expresion"])
            k = (niv, "Expresión", t)
            if k not in vistos[idi]:
                vistos[idi].add(k)
                datos[idi][niv].append(("Expresión", t, dict(r.get("traducciones", {})), sid))

    for p in sorted(base.joinpath("packs").glob("*.json")):
        d = json.loads(p.read_text(encoding="utf-8"))
        idi, niv = d["idioma"], d["nivel"]
        for v in d.get("vocabulario", []):
            t = limpio(v["item"])
            k = (niv, "Vocabulario", t)
            if k not in vistos[idi]:
                vistos[idi].add(k)
                datos[idi][niv].append(("Vocabulario", t, dict(v.get("traducciones", {})),
                                        d.get("topicId", "")))
    return datos


def hoja_idioma(wb, idioma, por_nivel):
    bases = [i for i in IDIOMAS if i != idioma]
    ws = wb.create_sheet(NOMBRE[idioma][:31])
    ws.freeze_panes = "A4"
    ws["A1"] = f"Traducciones — {NOMBRE[idioma]} (idioma que se aprende)"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Verde = ya existe, no tocar. Ámbar = falta traducir."
    ws["A2"].font = Font(size=9, color="7F6000")

    fila = 3
    cab = ["Nivel", "Tipo", f"{NOMBRE[idioma]} — original"] + [NOMBRE[b] for b in bases] + ["Origen"]
    for j, t in enumerate(cab, start=1):
        c = ws.cell(row=fila, column=j, value=t)
        c.font = Font(bold=True, color="FFFFFF", size=9); c.fill = CABEZA
        c.alignment = Alignment(wrap_text=True, horizontal="center")
    anchos = [8, 12, 34] + [22] * len(bases) + [12]
    for i, a_ in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = a_

    tot_celdas = 0
    for niv in ORDEN_NIVEL:
        filas = por_nivel.get(niv, [])
        if not filas:
            continue
        fila += 1
        c = ws.cell(row=fila, column=1, value=f"Nivel {niv} — {len(filas)} textos")
        c.font = Font(bold=True); c.fill = NIVEL_FILL
        for j in range(2, len(cab) + 1):
            ws.cell(row=fila, column=j).fill = NIVEL_FILL
        for tipo, texto, trads, origen in sorted(filas, key=lambda x: (x[0], x[1])):
            fila += 1
            ws.cell(row=fila, column=1, value=niv).border = BORDE
            ws.cell(row=fila, column=2, value=tipo).border = BORDE
            c = ws.cell(row=fila, column=3, value=texto)
            c.fill = GRIS; c.border = BORDE; c.alignment = Alignment(wrap_text=True)
            for j, b in enumerate(bases, start=4):
                cel = ws.cell(row=fila, column=j)
                if b in trads:
                    v = trads[b]
                    cel.value = v if v is not None else "—"
                    cel.fill = VERDE
                else:
                    cel.fill = AMBAR
                    tot_celdas += 1
                cel.border = BORDE
                cel.alignment = Alignment(wrap_text=True)
            ws.cell(row=fila, column=len(cab), value=origen).border = BORDE
    return tot_celdas


# --------------------------------------------------------------------------
# Prosa bilingue A2/B1
# --------------------------------------------------------------------------
def _valor(v):
    """Devuelve (texto_original, traducciones) para un campo bilingue.

    El campo puede ser un string plano (nadie lo tradujo todavia) o el objeto
    {idioma: texto} que admite el schema. En el objeto, el texto original es
    la clave del idioma que se aprende.
    """
    if isinstance(v, dict):
        return None, dict(v)
    return v, {}


def _campos_nucleo(d):
    """(ruta, valor) de cada campo bilingue de un nucleo, en orden de lectura."""
    yield "titulo", d.get("titulo")
    yield "descripcion", d.get("descripcion")
    c = d.get("cuadroReferencia") or {}
    if "titulo" in c:
        yield "cuadroReferencia.titulo", c["titulo"]
    for i, col in enumerate(c.get("columnas", [])):
        yield f"cuadroReferencia.columnas[{i}]", col
    for i, fila in enumerate(c.get("filas", [])):
        for j, celda in enumerate(fila):
            yield f"cuadroReferencia.filas[{i}][{j}]", celda
    if c.get("notaPie"):
        yield "cuadroReferencia.notaPie", c["notaPie"]
    for i, e in enumerate(d.get("ejemplos", [])):
        yield f"ejemplos[{i}].texto", e.get("texto")
    for clave in ("notas", "errores", "autochequeo"):
        for i, t in enumerate(d.get(clave, [])):
            yield f"{clave}[{i}]", t
    for i, r in enumerate(d.get("redemittel", [])):
        if "funcion" in r:
            yield f"redemittel[{i}].funcion", r["funcion"]
    yield "promptCorreccion", d.get("promptCorreccion")


def _campos_semana(d):
    """(ruta, valor) de cada campo de texto de una semana especial."""
    yield "titulo", d.get("titulo")
    yield "consigna", d.get("consigna")
    for clave in ("requisitos", "microtareas", "autochequeo"):
        for i, t in enumerate(d.get(clave, [])):
            yield f"{clave}[{i}]", t
    yield "promptCorreccion", d.get("promptCorreccion")


def _campos_aparicion(a):
    yield "subtitulo", a.get("subtitulo")
    m = a.get("mision") or {}
    if "consigna" in m:
        yield "mision.consigna", m["consigna"]
    for i, r in enumerate(m.get("requisitos", [])):
        yield f"mision.requisitos[{i}]", r
    for i, mt in enumerate(a.get("microtareas", [])):
        yield f"microtareas[{i}].texto", mt.get("texto")


def recolectar_prosa(base):
    """{(idioma, nivel): [(origen, ruta, texto, traducciones, otros_usos)]}.

    Los textos repetidos (microtareas y requisitos casi iguales entre
    apariciones, celdas de cuadro que se repiten) se agrupan en una sola fila:
    se traducen una vez y el importador los vuelca en todos los lugares donde
    aparece ese texto. `otros_usos` lista los demas lugares, solo informativo.
    """
    crudo = defaultdict(list)

    for p in sorted(base.joinpath("nucleos").glob("*.json")):
        d = json.loads(p.read_text(encoding="utf-8"))
        if d.get("_tipo") or d["nivel"] not in NIVELES_BILINGUES:
            continue
        origen = f'{d["skillId"]}-{d["nivel"]}'
        for ruta, v in _campos_nucleo(d):
            if v is None:
                continue
            texto, trads = _valor(v)
            crudo[(d["idioma"], d["nivel"])].append((origen, ruta, texto, trads))

    for p in sorted(base.joinpath("ocurrencias").glob("*.json")):
        d = json.loads(p.read_text(encoding="utf-8"))
        if d["nivel"] not in NIVELES_BILINGUES:
            continue
        for a in d["apariciones"]:
            origen = f'{d["idioma"]}-{d["nivel"]}-{d["anio"]} · {a["skillId"]} · {a.get("topicId","")} · ap{a.get("order","")}'
            for ruta, v in _campos_aparicion(a):
                if v is None:
                    continue
                texto, trads = _valor(v)
                crudo[(d["idioma"], d["nivel"])].append((origen, ruta, texto, trads))

    return _agrupar(crudo)


def recolectar_repaso(base):
    """Igual que recolectar_prosa, pero para las semanas especiales A2/B1."""
    crudo = defaultdict(list)
    for p in sorted(base.joinpath("nucleos").glob("*.json")):
        d = json.loads(p.read_text(encoding="utf-8"))
        if d.get("_tipo") != "semana_especial" or d.get("nivel") not in NIVELES_BILINGUES:
            continue
        origen = d["id"]
        for ruta, v in _campos_semana(d):
            if v is None:
                continue
            texto, trads = _valor(v)
            crudo[(d["idioma"], d["nivel"])].append((origen, ruta, texto, trads))
    return _agrupar(crudo)


def _agrupar(crudo):
    """Junta en una fila los textos repetidos dentro de la misma combinacion."""
    datos = {}
    for clave, filas in crudo.items():
        agrupado = {}
        orden = []
        for origen, ruta, texto, trads in filas:
            k = texto if texto is not None else trads.get(clave[0])
            if k not in agrupado:
                agrupado[k] = [origen, ruta, texto, dict(trads), []]
                orden.append(k)
            else:
                agrupado[k][3].update({c: v for c, v in trads.items() if c not in agrupado[k][3]})
                agrupado[k][4].append(f"{origen} · {ruta}")
        datos[clave] = [tuple(agrupado[k]) for k in orden]
    return datos


def hoja_prosa(wb, idioma, nivel, filas, prefijo="Prosa", rotulo="Prosa bilingüe",
               limites=True, nota_extra=""):
    bases = [i for i in IDIOMAS if i != idioma]
    ws = wb.create_sheet(f"{prefijo} {idioma.upper()} {nivel}"[:31])
    ws.freeze_panes = "A4"
    ws["A1"] = f"{rotulo} — {NOMBRE[idioma]} {nivel}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Verde = ya existe, no tocar. Ámbar = falta traducir. El texto original "
                f"queda como clave «{idioma}» del campo y lo pone el importador: no hay que copiarlo. "
                + ("«Largo» es el límite de caracteres que el schema le exige a cada traducción. "
                   if limites else "")
                + nota_extra)
    ws["A2"].font = Font(size=9, color="7F6000")

    cab = (["Origen", "Campo", "Largo", f"{NOMBRE[idioma]} — original"]
           + [NOMBRE[b] for b in bases] + ["Otros usos del mismo texto"])
    fila = 3
    for j, t in enumerate(cab, start=1):
        c = ws.cell(row=fila, column=j, value=t)
        c.font = Font(bold=True, color="FFFFFF", size=9); c.fill = CABEZA
        c.alignment = Alignment(wrap_text=True, horizontal="center")
    for i, a_ in enumerate([26, 24, 9, 46] + [30] * len(bases) + [40], start=1):
        ws.column_dimensions[get_column_letter(i)].width = a_

    tot = 0
    origen_actual = None
    for origen, ruta, texto, trads, otros in filas:
        if origen != origen_actual:
            origen_actual = origen
            fila += 1
            c = ws.cell(row=fila, column=1, value=origen)
            c.font = Font(bold=True); c.fill = NIVEL_FILL
            for j in range(2, len(cab) + 1):
                ws.cell(row=fila, column=j).fill = NIVEL_FILL
        fila += 1
        ws.cell(row=fila, column=1, value=origen).border = BORDE
        ws.cell(row=fila, column=2, value=ruta).border = BORDE
        c = ws.cell(row=fila, column=3, value=_limite(ruta) if limites else "")
        c.border = BORDE; c.font = Font(size=8, color="777777")
        c.alignment = Alignment(horizontal="center")
        c = ws.cell(row=fila, column=4, value=texto if texto is not None else trads.get(idioma, ""))
        c.fill = GRIS; c.border = BORDE; c.alignment = Alignment(wrap_text=True)
        for j, b in enumerate(bases, start=5):
            cel = ws.cell(row=fila, column=j)
            if b in trads:
                cel.value = trads[b]; cel.fill = VERDE
            else:
                cel.fill = AMBAR; tot += 1
            cel.border = BORDE; cel.alignment = Alignment(wrap_text=True)
        c = ws.cell(row=fila, column=len(cab),
                    value=f"+{len(otros)}: " + " | ".join(otros) if otros else "")
        c.border = BORDE; c.font = Font(size=8, color="777777")
        c.alignment = Alignment(wrap_text=True)
    return tot


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--contenido", default="../contenido")
    ap.add_argument("--salida", default="../traducciones-maestro.xlsx")
    a = ap.parse_args()
    datos = recolectar(Path(a.contenido))
    prosa = recolectar_prosa(Path(a.contenido))
    repaso = recolectar_repaso(Path(a.contenido))

    wb = openpyxl.Workbook()
    idx = wb.active; idx.title = "ÍNDICE"
    idx["A1"] = "Maestro de traducciones pendientes"
    idx["A1"].font = Font(bold=True, size=14)
    idx["A2"] = ("Dos bloques: una hoja por idioma que se aprende con vocabulario y "
                 "expresiones, y una hoja por combinación (idioma + nivel) con toda la prosa "
                 "de las fichas bilingües A2/B1. Rellenar solo las celdas en ámbar. "
                 "No incluye contraste ni errores contrastivos: eso no es traducción, es "
                 "análisis nuevo por cada par de idiomas y lo escribe el diseñador aparte.")
    idx["A2"].font = Font(size=9, color="555555")
    idx["A2"].alignment = Alignment(wrap_text=True)
    idx.column_dimensions["A"].width = 95

    for j, t in enumerate(["Idioma que se aprende", "Celdas por traducir"], start=1):
        c = idx.cell(row=4, column=j, value=t)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = CABEZA
    idx.column_dimensions["B"].width = 18

    fila = 5
    total = 0
    for idi in sorted(datos, key=lambda x: NOMBRE[x]):
        n = hoja_idioma(wb, idi, datos[idi])
        idx.cell(row=fila, column=1, value=NOMBRE[idi])
        idx.cell(row=fila, column=2, value=n)
        total += n
        fila += 1
    fila += 1
    c = idx.cell(row=fila, column=1, value="Prosa bilingüe A2/B1 (núcleos + apariciones)")
    c.font = Font(bold=True, color="FFFFFF"); c.fill = CABEZA
    idx.cell(row=fila, column=2).fill = CABEZA
    fila += 1
    for (idi, niv) in sorted(prosa, key=lambda x: (NOMBRE[x[0]], ORDEN_NIVEL.index(x[1]))):
        n = hoja_prosa(wb, idi, niv, prosa[(idi, niv)])
        idx.cell(row=fila, column=1, value=f"{NOMBRE[idi]} {niv} — {len(prosa[(idi, niv)])} textos")
        idx.cell(row=fila, column=2, value=n)
        total += n
        fila += 1
    fila += 1
    c = idx.cell(row=fila, column=1, value="Semanas de repaso A2/B1")
    c.font = Font(bold=True, color="FFFFFF"); c.fill = CABEZA
    idx.cell(row=fila, column=2).fill = CABEZA
    fila += 1
    for (idi, niv) in sorted(repaso, key=lambda x: (NOMBRE[x[0]], ORDEN_NIVEL.index(x[1]))):
        n = hoja_prosa(wb, idi, niv, repaso[(idi, niv)], prefijo="Repaso",
                       rotulo="Semanas de repaso", limites=False,
                       nota_extra="Estos campos no tienen límite de longitud. La palabra objetivo "
                                  "(vocabulario, conector, estructura) queda en el idioma que se "
                                  "aprende; el día, los minutos y el número de semana sí se traducen.")
        idx.cell(row=fila, column=1, value=f"{NOMBRE[idi]} {niv} — {len(repaso[(idi, niv)])} textos")
        idx.cell(row=fila, column=2, value=n)
        total += n
        fila += 1
    idx.cell(row=fila, column=1, value="TOTAL").font = Font(bold=True)
    idx.cell(row=fila, column=2, value=total).font = Font(bold=True)

    wb.save(a.salida)
    print(f"{total} celdas por traducir -> {a.salida}")


if __name__ == "__main__":
    main()
