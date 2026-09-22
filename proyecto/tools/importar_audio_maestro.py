#!/usr/bin/env python3
"""
importar_audio_maestro.py — lee las marcas de "Grabado" del Excel devuelto.

A diferencia de las traducciones, "grabado" no es un dato de contenido: no
tiene sentido escribirlo dentro de un nucleo o un pack. Vive en un archivo de
estado aparte (audio-estado.json), indexado por texto. El exportador lo lee y
pre-marca la columna la proxima vez.

Cada entrada del estado guarda tambien el NOMBRE DE ARCHIVO real que Fer le
dio al grabar (columna "Archivo audio" del Excel, si existe), no solo un
booleano. Eso es lo que permite despues generar la tabla de mapeo hacia
assets/audio/ sin tener que volver a preguntar.

Uso:
    python3 importar_audio_maestro.py ../audio-a-grabar.xlsx --estado ../audio-estado.json
"""
import argparse, json, re, sys
from pathlib import Path
import openpyxl

NOMBRE = {"de": "Alemán", "en": "Inglés", "es": "Español",
          "fr": "Francés", "it": "Italiano", "pt": "Portugués"}
COD = {v: k for k, v in NOMBRE.items()}


def limpio(t):
    return re.sub(r"\*+", "", str(t or "")).strip()


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser()
    ap.add_argument("excel")
    ap.add_argument("--estado", default="../audio-estado.json")
    a = ap.parse_args()

    estado_path = Path(a.estado)
    estado = json.loads(estado_path.read_text(encoding="utf-8")) if estado_path.exists() else {}

    wb = openpyxl.load_workbook(a.excel, data_only=True)
    marcados = 0
    for hoja in wb.sheetnames:
        if hoja.upper().startswith("ÍNDICE") or hoja.upper().startswith("INDICE"):
            continue
        idi = COD.get(hoja)
        if not idi:
            continue
        ws = wb[hoja]
        # La cabecera real esta en la fila 2 (la fila 3 es el primer separador de
        # nivel, "Nivel B1 - N textos"). Columna 6 "Archivo audio" es opcional.
        cab6 = ws.cell(row=2, column=6).value
        tiene_archivo = cab6 and "archivo" in str(cab6).lower()
        for r in range(4, ws.max_row + 1):
            tipo = ws.cell(row=r, column=2).value
            texto = ws.cell(row=r, column=3).value
            marca = ws.cell(row=r, column=5).value
            archivo = ws.cell(row=r, column=6).value if tiene_archivo else None
            if not tipo or not texto:
                continue
            if not marca and not archivo:
                continue
            k = f"{idi}|{tipo}|{limpio(texto)}"
            prev = estado.get(k)
            # migracion silenciosa del formato viejo (booleano) al nuevo (dict)
            entrada = prev if isinstance(prev, dict) else ({"grabado": True} if prev else {})
            entrada["grabado"] = bool(marca and str(marca).strip()) or entrada.get("grabado", False)
            if marca and str(marca).strip():
                # se guarda tal cual (p.ej. "Y (14.09.26)"): la fecha es dato de Fer
                entrada["marca"] = str(marca).strip()
            if archivo and str(archivo).strip():
                nuevo = str(archivo).strip()
                previo = entrada.get("archivo", "")
                # El estado se indexa sin nivel. Si el archivo es de OTRO nivel que
                # el principal (el mismo texto grabado en B2 y en C1), va a
                # otrosArchivos; si es del mismo nivel, es una correccion y reemplaza.
                if previo and previo != nuevo and previo[:6] != nuevo[:6]:
                    otros = entrada.setdefault("otrosArchivos", [])
                    if nuevo not in otros:
                        otros.append(nuevo)
                else:
                    entrada["archivo"] = nuevo
            estado[k] = entrada
            marcados += 1

    estado_path.write_text(json.dumps(estado, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{marcados} marcas de 'grabado' guardadas en {a.estado}")


if __name__ == "__main__":
    main()
